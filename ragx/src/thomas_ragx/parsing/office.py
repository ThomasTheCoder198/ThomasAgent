from collections.abc import Iterator, Sequence
from dataclasses import dataclass, field
from enum import StrEnum
from pathlib import Path
from typing import Any, Final

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.shapes.group import GroupShape

from thomas_ragx.parsing.model import (
    BBox,
    Block,
    BlockCollector,
    BlockType,
    ParsedDocument,
    ParseProvider,
    TableData,
    docx_paragraph,
    pdf_bbox,
)

HEADING_STYLE_PREFIX: Final = "Heading "
TITLE_STYLE: Final = "Title"
TITLE_LEVEL: Final = 1
EMU_PER_POINT: Final = 12700
# Word's w:outlineLvl runs 0-8 for headings 1-9; 9 means body text.
OUTLINE_BODY_LEVEL: Final = 9
MAX_STYLE_DEPTH: Final = 10
LIST_STYLE_PREFIXES: Final = ("List Bullet", "List Number")
ORDERED_STYLE_MARK: Final = "Number"
UNORDERED_FORMATS: Final = frozenset({"bullet", "none"})
NO_NUMBERING: Final = "0"
BULLET_MARKER: Final = "- "
LIST_INDENT: Final = "  "
NUMBER_MARKER_SUFFIX: Final = ". "
LINE_BREAK: Final = "\v"
# Footers, dates and slide numbers repeat on every slide and carry no slide content.
BOILERPLATE_PLACEHOLDERS: Final = frozenset(
    {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.HEADER}
)
# Body placeholders show bullets from their layout even when a paragraph has no bullet element.
BULLETED_PLACEHOLDERS: Final = frozenset({PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT})
BULLET_NONE: Final = "a:buNone"
BULLET_CHAR: Final = "a:buChar"
BULLET_NUMBER: Final = "a:buAutoNum"
OUTLINE_LEVEL_XPATH: Final = "./w:pPr/w:outlineLvl/@w:val"
NUM_ID_XPATH: Final = "./w:pPr/w:numPr/w:numId/@w:val"
NUM_LEVEL_XPATH: Final = "./w:pPr/w:numPr/w:ilvl/@w:val"
# numbering.xml tags and attributes, looked up with find/get so a hostile id never reaches an XPath string.
NUMBERING_NUM: Final = "w:num"
NUMBERING_NUM_ID: Final = "w:numId"
NUMBERING_ABSTRACT_ID: Final = "w:abstractNumId"
NUMBERING_ABSTRACT: Final = "w:abstractNum"
NUMBERING_LEVEL: Final = "w:lvl"
NUMBERING_LEVEL_INDEX: Final = "w:ilvl"
NUMBERING_FORMAT: Final = "w:numFmt"
VALUE_ATTRIBUTE: Final = "w:val"


def _outline_level(element: Any) -> int | None:
    values = element.xpath(OUTLINE_LEVEL_XPATH)
    return int(values[0]) if values else None


def _style_chain(paragraph: Paragraph) -> Iterator[Any]:
    style: Any = paragraph.style
    for _ in range(MAX_STYLE_DEPTH):  # bounded in case a damaged file has a style cycle
        if style is None:
            return
        yield style
        style = style.base_style


def _level_from_outline(outline: int) -> int | None:
    return outline + 1 if outline < OUTLINE_BODY_LEVEL else None


def _style_heading(style: Any) -> tuple[bool, int | None]:
    """(decided, level). An explicit outline level, body text (9) included, ends the walk up the style
    chain: Word's `TOC Heading` is based on Heading 1 but opts out of the outline."""
    outline = _outline_level(style.element)
    if outline is not None:
        return True, _level_from_outline(outline)
    name = style.name or ""
    if name == TITLE_STYLE:
        return True, TITLE_LEVEL
    suffix = name.removeprefix(HEADING_STYLE_PREFIX)
    if name.startswith(HEADING_STYLE_PREFIX) and suffix.isdigit():
        return True, int(suffix)
    return False, None


def _heading_level(paragraph: Paragraph) -> int | None:
    """Direct outline level first, then the style and its ancestors (a custom style based on Heading 2)."""
    direct = _outline_level(paragraph._p)  # python-docx has no public outline API
    if direct is not None:
        return _level_from_outline(direct)
    for style in _style_chain(paragraph):
        decided, level = _style_heading(style)
        if decided:
            return level
    return None


class _Marker(StrEnum):
    BULLET = "bullet"
    NUMBER = "number"


@dataclass(frozen=True)
class _ListItem:
    marker: _Marker
    depth: int


@dataclass
class _ListLines:
    """Renders list items as indented markdown-style lines, numbering each level on its own."""

    lines: list[str] = field(default_factory=list)
    counters: dict[int, int] = field(default_factory=dict)

    def add(self, depth: int, marker: _Marker | None, text: str) -> None:
        for deeper in [level for level in self.counters if level > depth]:
            del self.counters[deeper]
        self.counters[depth] = self.counters.get(depth, 0) + 1
        numbered = f"{self.counters[depth]}{NUMBER_MARKER_SUFFIX}"
        prefix = {None: "", _Marker.BULLET: BULLET_MARKER, _Marker.NUMBER: numbered}[marker]
        self.lines.append(LIST_INDENT * depth + prefix + text)


@dataclass
class _ListRun:
    """Consecutive list paragraphs become one block located at the first paragraph (ruling O6-3)."""

    first_paragraph: int
    top_marker: _Marker
    rendered: _ListLines = field(default_factory=_ListLines)


def _numbering_ids(element: Any) -> tuple[str | None, int | None]:
    num_id = element.xpath(NUM_ID_XPATH)
    level = element.xpath(NUM_LEVEL_XPATH)
    return (num_id[0] if num_id else None, int(level[0]) if level else None)


def _style_depth(name: str) -> int:
    """`List Bullet 2` is the second level when the style carries no w:ilvl."""
    suffix = name.rsplit(" ", 1)[-1]
    return int(suffix) - 1 if suffix.isdigit() else 0


def _abstract_id(root: Any, num_id: str) -> str | None:
    for num in root.findall(qn(NUMBERING_NUM)):
        if num.get(qn(NUMBERING_NUM_ID)) == num_id:
            abstract = num.find(qn(NUMBERING_ABSTRACT_ID))
            return None if abstract is None else abstract.get(qn(VALUE_ATTRIBUTE))
    return None


class _Numbering:
    """Looks up whether a numbering id is a decimal list or a bullet list in numbering.xml."""

    def __init__(self, document: DocxDocument) -> None:
        self._document = document

    def is_ordered(self, num_id: str, level: int) -> bool | None:
        try:
            root = self._document.part.numbering_part.element
        except NotImplementedError:  # a document without numbering.xml cannot have numbered lists
            return None
        abstract_id = _abstract_id(root, num_id)
        if abstract_id is None:
            return None
        for abstract in root.findall(qn(NUMBERING_ABSTRACT)):
            if abstract.get(qn(NUMBERING_ABSTRACT_ID)) != abstract_id:
                continue
            for level_node in abstract.findall(qn(NUMBERING_LEVEL)):
                if level_node.get(qn(NUMBERING_LEVEL_INDEX)) == str(level):
                    number_format = level_node.find(qn(NUMBERING_FORMAT))
                    value = None if number_format is None else number_format.get(qn(VALUE_ATTRIBUTE))
                    return None if value is None else value not in UNORDERED_FORMATS
        return None


def _list_item(paragraph: Paragraph, numbering: _Numbering) -> _ListItem | None:
    num_id, level = _numbering_ids(paragraph._p)  # python-docx has no public numbering API
    names = [style.name or "" for style in _style_chain(paragraph)]
    for style in _style_chain(paragraph):
        style_id, style_level = _numbering_ids(style.element)
        num_id = num_id or style_id
        level = level if level is not None else style_level
    list_style = next((name for name in names if name.startswith(LIST_STYLE_PREFIXES)), None)
    if num_id == NO_NUMBERING or (num_id is None and list_style is None):
        return None
    depth = level if level is not None else _style_depth(list_style or "")
    ordered = numbering.is_ordered(num_id, depth) if num_id is not None else None
    if ordered is None:
        ordered = ORDERED_STYLE_MARK in (list_style or "")
    return _ListItem(_Marker.NUMBER if ordered else _Marker.BULLET, depth)


def _docx_grid(table: Table) -> list[tuple[str, ...]]:
    """python-docx repeats a merged cell in every grid slot it covers; a row that starts late or ends early
    (w:gridBefore / w:gridAfter) is padded with empty cells, so every row has the same width."""
    rows = [
        ("",) * row.grid_cols_before
        + tuple(cell.text.strip() for cell in row.cells)
        + ("",) * row.grid_cols_after
        for row in table.rows
    ]
    width = max([len(table.columns), *(len(row) for row in rows)])
    return [row + ("",) * (width - len(row)) for row in rows]


def _docx_table(table: Table) -> TableData | None:
    grid = _docx_grid(table)
    if not any(any(row) for row in grid):
        return None
    return TableData(header=grid[0], rows=tuple(grid[1:]))


class _DocxReader:
    def __init__(self, document: DocxDocument) -> None:
        self._numbering = _Numbering(document)
        self._collector = BlockCollector()
        self._run: _ListRun | None = None

    def read(self, document: DocxDocument) -> tuple[Block, ...]:
        for index, item in enumerate(document.iter_inner_content()):
            if isinstance(item, Table):
                self._flush()
                table = _docx_table(item)
                if table is not None:
                    self._collector.add(BlockType.TABLE, "", docx_paragraph(index), table=table)
            else:
                self._paragraph(index, item)
        self._flush()
        return self._collector.blocks()

    def _paragraph(self, index: int, paragraph: Paragraph) -> None:
        level = _heading_level(paragraph)
        list_item = None if level is not None else _list_item(paragraph, self._numbering)
        text = paragraph.text.strip()
        if list_item is not None:
            if text:  # a blank list paragraph neither adds a line nor ends the list
                self._list_line(index, list_item, text)
            return
        self._flush()
        block_type = BlockType.HEADING if level is not None else BlockType.TEXT
        self._collector.add(block_type, text, docx_paragraph(index), level=level)

    def _list_line(self, index: int, item: _ListItem, text: str) -> None:
        if self._run is not None and item.depth == 0 and item.marker != self._run.top_marker:
            self._flush()  # an ordered list directly followed by a bullet list is two lists
        if self._run is None:
            self._run = _ListRun(index, item.marker)
        self._run.rendered.add(item.depth, item.marker, text)

    def _flush(self) -> None:
        if self._run is not None:
            self._collector.add(
                BlockType.LIST, "\n".join(self._run.rendered.lines), docx_paragraph(self._run.first_paragraph)
            )
            self._run = None


class DocxParser:
    """DOCX has no fixed layout: the locator is the 0-based body index (ruling A22).

    The index counts the paragraphs and tables that python-docx's `iter_inner_content` yields, so a
    viewer can resolve it with the same call. Text inside pictures is not read (no OCR, ruling O5).
    """

    def parse(self, path: Path) -> ParsedDocument:
        document = Document(str(path))
        return ParsedDocument(_DocxReader(document).read(document), ParseProvider.LOCAL)


def _shape_box(shape: Any) -> BBox | None:
    left, top = getattr(shape, "left", None), getattr(shape, "top", None)
    width, height = getattr(shape, "width", None), getattr(shape, "height", None)
    if left is None or top is None or width is None or height is None:
        return None  # an inherited placeholder position
    return (
        left / EMU_PER_POINT,
        top / EMU_PER_POINT,
        (left + width) / EMU_PER_POINT,
        (top + height) / EMU_PER_POINT,
    )


def _bullet_marker(paragraph: Any, inherits_bullets: bool) -> _Marker | None:
    properties = paragraph._p.pPr  # python-pptx has no public bullet API
    if properties is not None:
        if properties.xpath(f"./{BULLET_NUMBER}"):
            return _Marker.NUMBER
        if properties.xpath(f"./{BULLET_CHAR}"):
            return _Marker.BULLET
        if properties.xpath(f"./{BULLET_NONE}"):
            return None
    return _Marker.BULLET if inherits_bullets else None


def _paragraph_text(paragraph: Any) -> str:
    return str(paragraph.text).replace(LINE_BREAK, "\n").strip()


def _inherits_bullets(shape: Any) -> bool:
    return bool(shape.is_placeholder) and shape.placeholder_format.type in BULLETED_PLACEHOLDERS


def _text_block(shape: Any) -> tuple[BlockType, str]:
    inherits = _inherits_bullets(shape)
    items = [
        (paragraph, text, _bullet_marker(paragraph, inherits))
        for paragraph in shape.text_frame.paragraphs
        if (text := _paragraph_text(paragraph))
    ]
    if not any(marker for _, _, marker in items):
        return BlockType.TEXT, "\n".join(text for _, text, _ in items)
    rendered = _ListLines()
    for paragraph, text, marker in items:
        rendered.add(int(paragraph.level), marker, text)
    return BlockType.LIST, "\n".join(rendered.lines)


def _pptx_grid(table: Any) -> list[list[str]]:
    """Every grid slot a merged cell covers repeats the origin's text, so merged headers and row groups
    stay rectangular (same rule as the Datalab tables, ruling S9-5)."""
    grid = [["" for _ in table.columns] for _ in table.rows]
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            if cell.is_spanned:
                continue
            text = cell.text.strip()
            rows = range(row_index, row_index + (cell.span_height if cell.is_merge_origin else 1))
            columns = range(column_index, column_index + (cell.span_width if cell.is_merge_origin else 1))
            for covered_row in rows:
                for covered_column in columns:
                    grid[covered_row][covered_column] = text
    return grid


def _pptx_table(table: Any) -> TableData | None:
    grid = _pptx_grid(table)
    if not any(any(row) for row in grid):
        return None
    return TableData(header=tuple(grid[0]), rows=tuple(tuple(row) for row in grid[1:]))


def _walk(
    shapes: Sequence[Any], title_id: int | None, inherited: BBox | None
) -> Iterator[tuple[Any, BBox | None]]:
    """Reading order is title first, then top-to-bottom, left-to-right. Group children are placed in
    the group's coordinate space, so they take the outermost group's box (slide space) and never the box
    of a nested group, which is expressed in its parent's child space."""
    for shape in sorted(shapes, key=lambda s: (s.shape_id != title_id, s.top or 0, s.left or 0)):
        box = inherited if inherited is not None else _shape_box(shape)
        if isinstance(shape, GroupShape):
            yield from _walk(list(shape.shapes), None, box)
        else:
            yield shape, box


class PptxParser:
    """A slide has a fixed layout: page = slide number, bbox in points (ruling A22).

    Pictures carry no text for us (no OCR, ruling O5). Footer, date and slide-number placeholders are skipped.
    """

    def parse(self, path: Path) -> ParsedDocument:
        presentation = Presentation(str(path))
        collector = BlockCollector()
        for number, slide in enumerate(presentation.slides, start=1):
            title = slide.shapes.title
            title_id = title.shape_id if title is not None else None
            for shape, box in _walk(list(slide.shapes), title_id, None):
                self._shape(collector, number, shape, box, is_title=shape.shape_id == title_id)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                collector.add(BlockType.TEXT, slide.notes_slide.notes_text_frame.text, pdf_bbox(number, None))
        return ParsedDocument(collector.blocks(), ParseProvider.LOCAL, page_count=len(presentation.slides))

    @staticmethod
    def _shape(
        collector: BlockCollector, number: int, shape: Any, box: BBox | None, *, is_title: bool
    ) -> None:
        locator = pdf_bbox(number, box)
        if shape.has_table:
            table = _pptx_table(shape.table)
            if table is not None:
                collector.add(BlockType.TABLE, "", locator, table=table)
        elif shape.has_text_frame and not _is_boilerplate(shape):
            if is_title:
                collector.add(
                    BlockType.HEADING, _paragraph_text(shape.text_frame), locator, level=TITLE_LEVEL
                )
            else:
                block_type, text = _text_block(shape)
                collector.add(block_type, text, locator)


def _is_boilerplate(shape: Any) -> bool:
    return bool(shape.is_placeholder) and shape.placeholder_format.type in BOILERPLATE_PLACEHOLDERS
