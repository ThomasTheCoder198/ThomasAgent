import csv
from collections.abc import Iterable, Sequence
from pathlib import Path
from typing import Final

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from thomas_ragx.parsing.model import (
    BlockCollector,
    BlockType,
    Locator,
    LocatorKind,
    ParsedDocument,
    ParseProvider,
    TableData,
    sheet_range,
)
from thomas_ragx.parsing.text_files import read_text

# A CSV has no sheets. The pipeline parses a temp file named `<file_id>-<version>.csv`, so the stem would
# leak an id into every citation; a fixed name keeps the locator stable and the same for every CSV.
CSV_SHEET_NAME: Final = "csv"
FIRST_COLUMN: Final = "A"
SNIFF_BYTES: Final = 4096
HEADING_LEVEL_SHEET: Final = 1


def _cell(value: object) -> str:
    return "" if value is None else str(value).strip()


def table_block(collector: BlockCollector, sheet: str, rows: Iterable[Sequence[object]]) -> None:
    """Empty rows are skipped but keep their numbers, so a cited range matches what the user sees."""
    numbered = [(number, tuple(_cell(v) for v in row)) for number, row in enumerate(rows, start=1)]
    kept = [(number, cells) for number, cells in numbered if any(cells)]
    if not kept:
        return
    # Trailing empty columns (styled but blank cells stretch an xlsx sheet) are not part of the table.
    width = max(max(index for index, value in enumerate(cells, start=1) if value) for _, cells in kept)
    padded = [(number, (cells + ("",) * width)[:width]) for number, cells in kept]
    (header_row, header), body = padded[0], padded[1:]
    last_column = get_column_letter(width)
    table = TableData(
        header=header,
        rows=tuple(cells for _, cells in body),
        source_rows=tuple(number for number, _ in body),
        columns=(FIRST_COLUMN, last_column),
    )
    last_row = body[-1][0] if body else header_row
    locator = sheet_range(sheet, f"{FIRST_COLUMN}{header_row}:{last_column}{last_row}")
    collector.add(BlockType.TABLE, "", locator, table=table)


class CsvParser:
    def parse(self, path: Path) -> ParsedDocument:
        text = read_text(path)
        try:
            dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(text[:SNIFF_BYTES])
        except csv.Error:
            dialect = csv.excel
        collector = BlockCollector()
        table_block(collector, CSV_SHEET_NAME, csv.reader(text.splitlines(), dialect))
        return ParsedDocument(collector.blocks(), ParseProvider.LOCAL)


class XlsxParser:
    """data_only: cached formula results; a formula never calculated by Excel reads as empty (openpyxl)."""

    def parse(self, path: Path) -> ParsedDocument:
        book = load_workbook(path, read_only=True, data_only=True)
        collector = BlockCollector()
        try:
            for sheet in book.worksheets:
                rows = list(sheet.iter_rows(values_only=True))
                if any(any(_cell(v) for v in row) for row in rows):
                    # The heading names the sheet only, so its locator has no range.
                    heading = Locator(LocatorKind.SHEET_RANGE, sheet=sheet.title)
                    collector.add(BlockType.HEADING, sheet.title, heading, level=HEADING_LEVEL_SHEET)
                    table_block(collector, sheet.title, rows)
        finally:
            book.close()
        return ParsedDocument(collector.blocks(), ParseProvider.LOCAL)
