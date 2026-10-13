from pathlib import Path
from typing import Any

import pymupdf
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.oxml.ns import qn as pptx_qn
from pptx.util import Inches

from thomas_ragx.parsing.model import BlockType, LocatorKind
from thomas_ragx.parsing.office import EMU_PER_POINT, DocxParser, PptxParser


def _png(path: Path) -> Path:
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 8, 8), False)
    path.write_bytes(pixmap.tobytes("png"))
    return path


def test_docx_parser_keeps_body_order_headings_tables_and_paragraph_index(tmp_path: Path) -> None:
    document = Document()
    document.add_heading("Điều 1", level=1)
    document.add_paragraph("")
    document.add_paragraph("Nội dung điều một.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Mục", "Giá"
    table.cell(1, 0).text, table.cell(1, 1).text = "A", "5"
    path = tmp_path / "a.docx"
    document.save(path)
    blocks = DocxParser().parse(path).blocks
    assert [(b.block_type, b.locator.paragraph) for b in blocks] == [
        (BlockType.HEADING, 0),
        (BlockType.TEXT, 2),
        (BlockType.TABLE, 3),
    ]
    assert blocks[0].level == 1 and {b.locator.kind for b in blocks} == {LocatorKind.DOCX_PARAGRAPH}


def test_docx_headings_come_from_word_styles_their_ancestors_and_outline_levels(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Tên tài liệu", style="Title")
    document.add_heading("Mục ba", level=3)
    custom = document.styles.add_style("Chương", WD_STYLE_TYPE.PARAGRAPH)
    custom.base_style = document.styles["Heading 2"]
    document.add_paragraph("Chương một", style="Chương")
    outlined = document.add_paragraph("Đề mục thủ công")
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "3")
    outlined._p.get_or_add_pPr().append(outline)
    document.add_paragraph("Đoạn thường")
    path = tmp_path / "h.docx"
    document.save(path)
    blocks = DocxParser().parse(path).blocks
    assert [(b.block_type, b.level, b.text) for b in blocks] == [
        (BlockType.HEADING, 1, "Tên tài liệu"),
        (BlockType.HEADING, 3, "Mục ba"),
        (BlockType.HEADING, 2, "Chương một"),
        (BlockType.HEADING, 4, "Đề mục thủ công"),
        (BlockType.TEXT, None, "Đoạn thường"),
    ]


def test_docx_consecutive_list_items_become_one_list_block_at_the_first_paragraph(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Mở đầu")
    document.add_paragraph("một", style="List Bullet")
    document.add_paragraph("hai", style="List Bullet")
    document.add_paragraph("hai-a", style="List Bullet 2")
    document.add_paragraph("Giữa")
    document.add_paragraph("x", style="List Number")
    document.add_paragraph("y", style="List Number")
    path = tmp_path / "l.docx"
    document.save(path)
    blocks = DocxParser().parse(path).blocks
    assert [(b.block_type, b.text, b.locator.paragraph) for b in blocks] == [
        (BlockType.TEXT, "Mở đầu", 0),
        (BlockType.LIST, "- một\n- hai\n  - hai-a", 1),
        (BlockType.TEXT, "Giữa", 4),
        (BlockType.LIST, "1. x\n2. y", 5),
    ]


def test_docx_direct_numbering_on_list_paragraph_uses_the_numbering_part_for_kind_and_depth(
    tmp_path: Path,
) -> None:
    document = Document()
    # numId 7 is decimal and numId 1 is a bullet in python-docx's default numbering part.
    for text, num_id, level in (("một", "7", "0"), ("một-a", "1", "1"), ("hai", "7", "0")):
        paragraph = document.add_paragraph(text, style="List Paragraph")
        numbering = OxmlElement("w:numPr")
        for tag, value in (("w:ilvl", level), ("w:numId", num_id)):
            child = OxmlElement(tag)
            child.set(qn("w:val"), value)
            numbering.append(child)
        paragraph._p.get_or_add_pPr().append(numbering)
    path = tmp_path / "n.docx"
    document.save(path)
    (block,) = DocxParser().parse(path).blocks
    assert (block.block_type, block.text, block.locator.paragraph) == (
        BlockType.LIST,
        "1. một\n  - một-a\n2. hai",
        0,
    )


def test_docx_merged_table_cells_become_a_rectangular_grid(tmp_path: Path) -> None:
    document = Document()
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Giá"
    table.cell(0, 2).text = "Ghi chú"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "Nhóm A"
    table.cell(1, 1).text, table.cell(1, 2).text = "1", "x"
    table.cell(2, 1).text, table.cell(2, 2).text = "2", "y"
    path = tmp_path / "m.docx"
    document.save(path)
    (block,) = DocxParser().parse(path).blocks
    assert block.table is not None
    assert block.table.header == ("Giá", "Giá", "Ghi chú")
    assert block.table.rows == (("Nhóm A", "1", "x"), ("Nhóm A", "2", "y"))


def test_docx_blank_table_and_picture_only_paragraphs_yield_no_text(tmp_path: Path) -> None:
    document = Document()
    document.add_table(rows=2, cols=2)
    document.add_picture(str(_png(tmp_path / "p.png")))  # no OCR on the server: a picture has no text (O5)
    path = tmp_path / "pic.docx"
    document.save(path)
    assert DocxParser().parse(path).blocks == ()


def test_pptx_parser_uses_slide_numbers_and_point_boxes(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])  # title only
    slide.shapes.title.text = "Doanh thu"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1))
    box.text_frame.text = "Tăng 20%"
    path = tmp_path / "a.pptx"
    presentation.save(path)
    blocks = PptxParser().parse(path).blocks
    assert [(b.block_type, b.text, b.locator.page) for b in blocks] == [
        (BlockType.HEADING, "Doanh thu", 1),
        (BlockType.TEXT, "Tăng 20%", 1),
    ]
    assert blocks[1].locator.bbox == (
        Inches(1) / EMU_PER_POINT,
        Inches(2) / EMU_PER_POINT,
        Inches(4) / EMU_PER_POINT,
        Inches(3) / EMU_PER_POINT,
    )


def test_pptx_title_comes_first_and_body_placeholder_is_a_nested_list(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])  # title and content
    slide.shapes.title.text = "Kế hoạch"
    body = slide.placeholders[1].text_frame
    body.text = "Quý 1"
    nested = body.add_paragraph()
    nested.text = "Tuyển người"
    nested.level = 1
    body.add_paragraph().text = "Quý 2"
    top = slide.shapes.add_textbox(0, 0, Inches(2), Inches(0.2))
    top.text_frame.text = "Nhãn trên cùng"
    path = tmp_path / "s.pptx"
    presentation.save(path)
    blocks = PptxParser().parse(path).blocks
    assert [(b.block_type, b.level, b.text) for b in blocks] == [
        (BlockType.HEADING, 1, "Kế hoạch"),
        (BlockType.TEXT, None, "Nhãn trên cùng"),
        (BlockType.LIST, None, "- Quý 1\n  - Tuyển người\n- Quý 2"),
    ]


def test_pptx_textbox_is_a_list_only_when_its_paragraphs_have_bullets(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])  # blank
    plain = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    plain.text_frame.text = "dòng 1"
    plain.text_frame.add_paragraph().text = "dòng 2"
    bulleted = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(3), Inches(1))
    bulleted.text_frame.text = "mục a"
    bulleted.text_frame.add_paragraph().text = "mục b"
    for paragraph in bulleted.text_frame.paragraphs:
        properties = paragraph._p.get_or_add_pPr()
        properties.append(properties.makeelement(pptx_qn("a:buChar"), {"char": "-"}))
    path = tmp_path / "b.pptx"
    presentation.save(path)
    assert [(b.block_type, b.text) for b in PptxParser().parse(path).blocks] == [
        (BlockType.TEXT, "dòng 1\ndòng 2"),
        (BlockType.LIST, "- mục a\n- mục b"),
    ]


def test_pptx_auto_numbered_paragraphs_are_a_numbered_list(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    box.text_frame.text = "bước một"
    box.text_frame.add_paragraph().text = "bước hai"
    for paragraph in box.text_frame.paragraphs:
        properties = paragraph._p.get_or_add_pPr()
        properties.append(properties.makeelement(pptx_qn("a:buAutoNum"), {"type": "arabicPeriod"}))
    path = tmp_path / "n.pptx"
    presentation.save(path)
    (block,) = PptxParser().parse(path).blocks
    assert (block.block_type, block.text) == (BlockType.LIST, "1. bước một\n2. bước hai")


def test_pptx_merged_table_cells_become_a_rectangular_grid_with_a_frame_box(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = slide.shapes.add_table(3, 3, Inches(1), Inches(1), Inches(6), Inches(2))
    table = frame.table
    table.cell(0, 0).text, table.cell(0, 2).text = "Giá", "Ghi chú"
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 0).text = "Nhóm A"
    table.cell(1, 1).text, table.cell(1, 2).text = "1", "x"
    table.cell(2, 1).text, table.cell(2, 2).text = "2", "y"
    table.cell(1, 0).merge(table.cell(2, 0))
    path = tmp_path / "t.pptx"
    presentation.save(path)
    (block,) = PptxParser().parse(path).blocks
    assert block.table is not None
    assert block.table.header == ("Giá", "Giá", "Ghi chú")
    assert block.table.rows == (("Nhóm A", "1", "x"), ("Nhóm A", "2", "y"))
    assert block.locator.page == 1 and block.locator.bbox is not None
    assert block.locator.bbox[0] == Inches(1) / EMU_PER_POINT


def test_pptx_reads_grouped_shapes_and_notes_and_ignores_pictures(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text_frame.text = "Trong nhóm"
    slide.shapes.add_picture(str(_png(tmp_path / "p.png")), Inches(5), Inches(5))
    slide.notes_slide.notes_text_frame.text = "Ghi chú diễn giả"
    path = tmp_path / "g.pptx"
    presentation.save(path)
    blocks = PptxParser().parse(path).blocks
    assert [(b.text, b.locator.page) for b in blocks] == [("Trong nhóm", 1), ("Ghi chú diễn giả", 1)]
    assert blocks[0].locator.bbox == (
        group.left / EMU_PER_POINT,
        group.top / EMU_PER_POINT,
        (group.left + group.width) / EMU_PER_POINT,
        (group.top + group.height) / EMU_PER_POINT,
    )
    assert blocks[1].locator.bbox is None


def test_pptx_with_only_a_picture_yields_no_blocks(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(str(_png(tmp_path / "p.png")), Inches(1), Inches(1))
    path = tmp_path / "pic.pptx"
    presentation.save(path)
    assert PptxParser().parse(path).blocks == ()


def test_docx_style_that_opts_out_of_the_outline_is_body_text_even_when_based_on_a_heading(
    tmp_path: Path,
) -> None:
    document = Document()
    toc = document.styles.add_style("Contents Title", WD_STYLE_TYPE.PARAGRAPH)
    toc.base_style = document.styles["Heading 1"]
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "9")
    toc.element.get_or_add_pPr().append(outline)
    document.add_paragraph("Mục lục", style="Contents Title")
    document.add_heading("Chương 1", level=1)
    path = tmp_path / "toc.docx"
    document.save(path)
    assert [(b.block_type, b.level, b.text) for b in DocxParser().parse(path).blocks] == [
        (BlockType.TEXT, None, "Mục lục"),
        (BlockType.HEADING, 1, "Chương 1"),
    ]


def _trim_row(row: Any, *, before: int = 0, after: int = 0) -> None:
    """Make a row start late or end early the way Word does, with w:gridBefore / w:gridAfter."""
    tr = row._tr
    properties = tr.get_or_add_trPr()
    cells = tr.findall(qn("w:tc"))
    for tag, count, dropped in (
        ("w:gridBefore", before, cells[:before]),
        ("w:gridAfter", after, cells[len(cells) - after :] if after else []),
    ):
        if count:
            element = OxmlElement(tag)
            element.set(qn("w:val"), str(count))
            properties.append(element)
            for cell in dropped:
                tr.remove(cell)


def test_docx_rows_with_grid_before_or_after_are_padded_to_a_rectangular_grid(tmp_path: Path) -> None:
    document = Document()
    table = document.add_table(rows=3, cols=3)
    for row in range(3):
        for column in range(3):
            table.cell(row, column).text = f"{row}{column}"
    _trim_row(table.rows[1], before=1)
    _trim_row(table.rows[2], after=1)
    path = tmp_path / "ragged.docx"
    document.save(path)
    (block,) = DocxParser().parse(path).blocks
    assert block.table is not None
    assert block.table.header == ("00", "01", "02")
    assert block.table.rows == (("", "11", "12"), ("20", "21", ""))


def test_docx_blank_list_paragraph_does_not_split_the_list(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("a", style="List Bullet")
    document.add_paragraph("", style="List Bullet")
    document.add_paragraph("b", style="List Bullet")
    path = tmp_path / "blank-item.docx"
    document.save(path)
    (block,) = DocxParser().parse(path).blocks
    assert (block.block_type, block.text, block.locator.paragraph) == (BlockType.LIST, "- a\n- b", 0)


def test_docx_numbering_id_with_a_quote_is_data_not_xpath(tmp_path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph("mục", style="List Paragraph")
    numbering = OxmlElement("w:numPr")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "7'")
    numbering.append(num_id)
    paragraph._p.get_or_add_pPr().append(numbering)
    path = tmp_path / "hostile.docx"
    document.save(path)
    (block,) = DocxParser().parse(path).blocks
    assert (block.block_type, block.text) == (BlockType.LIST, "- mục")


def test_pptx_nested_group_children_take_the_outermost_group_box(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    outer = slide.shapes.add_group_shape()
    outer.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1)).text_frame.text = "Góc trên"
    inner = outer.shapes.add_group_shape()
    inner.shapes.add_textbox(Inches(6), Inches(4), Inches(1), Inches(1)).text_frame.text = "Nhóm con"
    path = tmp_path / "nested.pptx"
    presentation.save(path)
    blocks = PptxParser().parse(path).blocks
    outer_box = (
        outer.left / EMU_PER_POINT,
        outer.top / EMU_PER_POINT,
        (outer.left + outer.width) / EMU_PER_POINT,
        (outer.top + outer.height) / EMU_PER_POINT,
    )
    assert [b.text for b in blocks] == ["Góc trên", "Nhóm con"]
    assert {b.locator.bbox for b in blocks} == {outer_box}
