from pathlib import Path
from typing import Any

import pymupdf
import pytest
import structlog
import structlog.testing
from docx import Document
from opentelemetry.sdk.trace import ReadableSpan, TracerProvider
from opentelemetry.sdk.trace.export import SimpleSpanProcessor
from opentelemetry.sdk.trace.export.in_memory_span_exporter import InMemorySpanExporter
from opentelemetry.trace import StatusCode
from pptx import Presentation

from tests.unit.parsing.fixtures import datalab_tree
from tests.unit.parsing.test_pdf_local import _pdf
from thomas_ragx.parsing.datalab import ConvertOutcome
from thomas_ragx.parsing.model import LocatorKind, ParsedDocument, ParseProvider
from thomas_ragx.parsing.registry import ParserRegistry, local_registry
from thomas_ragx.parsing.service import SPAN_DATALAB, SPAN_LOCAL, ParseService
from thomas_ragx.platform.error_codes_gen import ErrorCode
from thomas_ragx.platform.errors import AppError

LOCAL_FIRST = frozenset({".docx", ".pptx"})


class FakeConverter:
    def __init__(
        self,
        qualities: list[float | None] | None = None,
        fail: bool = False,
        tree: dict[str, Any] | None = None,
        cost: dict[str, Any] | None = None,
        error: Exception | None = None,
    ) -> None:
        self.qualities = qualities or [4.0]
        self.fail = fail
        self.tree = tree
        self.cost = cost
        self.error = error
        self.modes: list[str] = []

    async def convert(self, path: Path, mode: str) -> ConvertOutcome:
        self.modes.append(mode)
        if self.error is not None:
            raise self.error
        if self.fail or len(self.modes) > len(self.qualities):
            raise TimeoutError
        tree = datalab_tree() if self.tree is None else self.tree
        return ConvertOutcome(tree, 2, self.qualities[len(self.modes) - 1], 1.0, self.cost)


def _service(
    converter: FakeConverter | None,
    local_first: frozenset[str] = LOCAL_FIRST,
    provider: TracerProvider | None = None,
) -> ParseService:
    return ParseService(
        local_registry(),
        converter,
        mode="balanced",
        retry_mode="accurate",
        min_quality=3.0,
        local_first_extensions=local_first,
        tracer=(provider or TracerProvider()).get_tracer(__name__),
    )


@pytest.fixture
def pdf(tmp_path: Path) -> Path:
    path = tmp_path / "a.pdf"
    _pdf(path)
    return path


@pytest.fixture
def docx(tmp_path: Path) -> Path:
    path = tmp_path / "a.docx"
    document = Document()
    document.add_heading("Điều 1", level=1)
    document.add_paragraph("Nội dung hợp đồng.")
    document.save(path)
    return path


@pytest.fixture
def empty_docx(tmp_path: Path) -> Path:
    path = tmp_path / "empty.docx"
    Document().save(path)
    return path


@pytest.fixture
def pptx(tmp_path: Path) -> Path:
    path = tmp_path / "a.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Tiêu đề"
    slide.placeholders[1].text = "Nội dung"
    deck.save(path)
    return path


async def test_parse_policy_uses_datalab_for_pdf_when_allowed(pdf: Path) -> None:
    converter = FakeConverter()
    parsed = await _service(converter).parse(pdf, ".pdf", external_parsing=True)
    assert (parsed.provider, parsed.degraded, converter.modes) == (ParseProvider.DATALAB, False, ["balanced"])


async def test_parse_policy_converts_datalab_boxes_to_points_using_the_pdf_page_size(tmp_path: Path) -> None:
    path = tmp_path / "a4.pdf"
    document = pymupdf.open()
    document.new_page(width=595, height=842)
    document.new_page(width=595, height=842)
    document.save(path)
    parsed = await _service(FakeConverter()).parse(path, ".pdf", external_parsing=True)
    bbox = parsed.blocks[1].locator.bbox
    assert bbox is not None and bbox[2] == pytest.approx(1397 * 595 / 1596, abs=0.01)


async def test_parse_policy_keeps_page_only_locators_when_the_source_size_is_unreadable(
    tmp_path: Path,
) -> None:
    path = tmp_path / "odd.pdf"
    path.write_bytes(b"%PDF-1.4 garbage")  # the stub converter does not care; the size reader does
    parsed = await _service(FakeConverter()).parse(path, ".pdf", external_parsing=True)
    assert all(b.locator.bbox is None and b.locator.page is not None for b in parsed.blocks)


async def test_parse_policy_retries_accurate_when_quality_is_low_and_keeps_the_better(pdf: Path) -> None:
    converter = FakeConverter(qualities=[2.0, 3.5])
    parsed = await _service(converter).parse(pdf, ".pdf", external_parsing=True)
    assert converter.modes == ["balanced", "accurate"] and parsed.quality == 3.5


async def test_parse_policy_keeps_the_first_result_when_the_retry_is_not_better(pdf: Path) -> None:
    converter = FakeConverter(qualities=[2.0, 1.0])
    parsed = await _service(converter).parse(pdf, ".pdf", external_parsing=True)
    assert converter.modes == ["balanced", "accurate"] and parsed.quality == 2.0


async def test_parse_policy_keeps_the_first_result_when_the_retry_fails(pdf: Path) -> None:
    converter = FakeConverter(qualities=[2.0])  # the second call raises
    parsed = await _service(converter).parse(pdf, ".pdf", external_parsing=True)
    assert (parsed.provider, parsed.quality, parsed.degraded) == (ParseProvider.DATALAB, 2.0, False)


async def test_parse_policy_does_not_retry_an_unknown_quality_and_stores_none(pdf: Path) -> None:
    converter = FakeConverter(qualities=[None])
    parsed = await _service(converter).parse(pdf, ".pdf", external_parsing=True)
    assert converter.modes == ["balanced"] and parsed.quality is None


async def test_parse_policy_does_not_retry_when_quality_meets_the_minimum(pdf: Path) -> None:
    converter = FakeConverter(qualities=[3.0])
    await _service(converter).parse(pdf, ".pdf", external_parsing=True)
    assert converter.modes == ["balanced"]


async def test_parse_policy_falls_back_locally_and_marks_degraded_when_datalab_fails(pdf: Path) -> None:
    parsed = await _service(FakeConverter(fail=True)).parse(pdf, ".pdf", external_parsing=True)
    assert (parsed.provider, parsed.degraded) == (ParseProvider.LOCAL, True)


async def test_parse_policy_never_calls_datalab_when_the_org_opted_out(pdf: Path) -> None:
    converter = FakeConverter()
    parsed = await _service(converter).parse(pdf, ".pdf", external_parsing=False)
    assert converter.modes == [] and parsed.degraded


async def test_parse_policy_without_a_key_parses_locally_degraded(pdf: Path) -> None:
    parsed = await _service(None).parse(pdf, ".pdf", external_parsing=True)
    assert (parsed.provider, parsed.degraded) == (ParseProvider.LOCAL, True)


async def test_parse_policy_parses_markdown_locally_without_degrading(tmp_path: Path) -> None:
    path = tmp_path / "a.md"
    path.write_text("# A\n\nB\n", encoding="utf-8")
    converter = FakeConverter()
    parsed = await _service(converter).parse(path, ".md", external_parsing=True)
    assert (parsed.provider, parsed.degraded, converter.modes) == (ParseProvider.LOCAL, False, [])


async def test_parse_policy_maps_a_corrupt_file_to_parse_failed(tmp_path: Path) -> None:
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"not a zip")
    with pytest.raises(AppError) as caught:
        await _service(None).parse(path, ".xlsx", external_parsing=True)
    assert caught.value.code is ErrorCode.INGEST_PARSE_FAILED


async def test_parse_policy_image_without_datalab_has_no_text(tmp_path: Path) -> None:
    path = tmp_path / "a.png"
    path.write_bytes(b"\x89PNG\r\n\x1a\n")
    with pytest.raises(AppError) as caught:
        await _service(None).parse(path, ".png", external_parsing=True)
    assert caught.value.code is ErrorCode.INGEST_NO_TEXT


async def test_parse_policy_parses_docx_locally_first_without_calling_datalab(docx: Path) -> None:
    converter = FakeConverter()
    parsed = await _service(converter).parse(docx, ".docx", external_parsing=True)
    assert (parsed.provider, parsed.degraded, converter.modes) == (ParseProvider.LOCAL, False, [])
    assert {b.locator.kind for b in parsed.blocks} == {LocatorKind.DOCX_PARAGRAPH}


async def test_parse_policy_parses_pptx_locally_first_without_calling_datalab(pptx: Path) -> None:
    converter = FakeConverter()
    parsed = await _service(converter).parse(pptx, ".pptx", external_parsing=True)
    assert (parsed.provider, parsed.degraded, converter.modes) == (ParseProvider.LOCAL, False, [])
    assert {b.locator.kind for b in parsed.blocks} == {LocatorKind.PDF_BBOX}


async def test_parse_policy_rescues_a_docx_without_blocks_with_datalab_not_degraded(empty_docx: Path) -> None:
    converter = FakeConverter()
    parsed = await _service(converter).parse(empty_docx, ".docx", external_parsing=True)
    assert (parsed.provider, parsed.degraded, converter.modes) == (ParseProvider.DATALAB, False, ["balanced"])
    assert {b.locator.kind for b in parsed.blocks} == {LocatorKind.PDF_BBOX}  # S9-1: boxes, not paragraphs


async def test_parse_policy_rescues_a_corrupt_docx_with_datalab(tmp_path: Path) -> None:
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a zip")
    parsed = await _service(FakeConverter()).parse(path, ".docx", external_parsing=True)
    assert (parsed.provider, parsed.degraded) == (ParseProvider.DATALAB, False)


async def test_parse_policy_corrupt_docx_without_a_key_is_parse_failed(tmp_path: Path) -> None:
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a zip")
    with pytest.raises(AppError) as caught:
        await _service(None).parse(path, ".docx", external_parsing=True)
    assert caught.value.code is ErrorCode.INGEST_PARSE_FAILED


async def test_parse_policy_empty_docx_without_a_key_has_no_text(empty_docx: Path) -> None:
    with pytest.raises(AppError) as caught:
        await _service(None).parse(empty_docx, ".docx", external_parsing=True)
    assert caught.value.code is ErrorCode.INGEST_NO_TEXT


async def test_parse_policy_never_rescues_when_the_org_opted_out(empty_docx: Path) -> None:
    converter = FakeConverter()
    with pytest.raises(AppError) as caught:
        await _service(converter).parse(empty_docx, ".docx", external_parsing=False)
    assert caught.value.code is ErrorCode.INGEST_NO_TEXT and converter.modes == []


async def test_parse_policy_keeps_the_local_error_when_the_rescue_fails_too(tmp_path: Path) -> None:
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a zip")
    with pytest.raises(AppError) as caught:
        await _service(FakeConverter(fail=True)).parse(path, ".docx", external_parsing=True)
    assert caught.value.code is ErrorCode.INGEST_PARSE_FAILED


async def test_parse_policy_local_first_set_is_configuration_not_code(docx: Path) -> None:
    converter = FakeConverter()
    parsed = await _service(converter, local_first=frozenset()).parse(docx, ".docx", external_parsing=True)
    assert (parsed.provider, converter.modes) == (ParseProvider.DATALAB, ["balanced"])


async def test_parse_policy_pdf_stays_datalab_first_with_the_default_local_first_set(pdf: Path) -> None:
    converter = FakeConverter()
    parsed = await _service(converter).parse(pdf, ".pdf", external_parsing=True)
    assert parsed.provider is ParseProvider.DATALAB and converter.modes == ["balanced"]


def _spans(exporter: InMemorySpanExporter) -> dict[str, ReadableSpan]:
    return {span.name: span for span in exporter.get_finished_spans()}


async def test_parse_spans_record_mode_and_leave_out_an_unknown_quality(pdf: Path) -> None:
    exporter = InMemorySpanExporter()
    provider = TracerProvider()
    provider.add_span_processor(SimpleSpanProcessor(exporter))
    await _service(FakeConverter(qualities=[None]), provider=provider).parse(
        pdf, ".pdf", external_parsing=True
    )
    attributes = _spans(exporter)[SPAN_DATALAB].attributes or {}
    assert attributes["parse.mode"] == "balanced" and attributes["parse.page_count"] == 2
    assert "parse.quality" not in attributes


async def test_parse_spans_mark_local_parsing(docx: Path) -> None:
    exporter = InMemorySpanExporter()
    provider = TracerProvider()
    provider.add_span_processor(SimpleSpanProcessor(exporter))
    await _service(None, provider=provider).parse(docx, ".docx", external_parsing=True)
    assert (_spans(exporter)[SPAN_LOCAL].attributes or {})["parse.provider"] == "local"


@pytest.fixture
def blank_pdf(tmp_path: Path) -> Path:
    path = tmp_path / "blank.pdf"
    document = pymupdf.open()
    document.new_page()
    document.save(path)
    return path


class FakeApiError(Exception):
    """Shaped like datalab_sdk.DatalabAPIError: a message that may hold anything, and a status code."""

    def __init__(self, status_code: int | None) -> None:
        super().__init__("secret-looking text")
        self.status_code = status_code


async def test_a_mapping_bug_is_a_parse_failure_not_a_datalab_outage(pdf: Path) -> None:
    converter = FakeConverter(tree={"children": 5})  # to_parsed_document cannot read this
    with pytest.raises(AppError) as caught:
        await _service(converter).parse(pdf, ".pdf", external_parsing=True)
    assert caught.value.code is ErrorCode.INGEST_PARSE_FAILED
    assert converter.modes == ["balanced"]  # no local fallback ran


async def test_a_mapping_bug_in_a_docx_rescue_is_a_parse_failure(empty_docx: Path) -> None:
    with pytest.raises(AppError) as caught:
        await _service(FakeConverter(tree={"children": 5})).parse(empty_docx, ".docx", external_parsing=True)
    assert caught.value.code is ErrorCode.INGEST_PARSE_FAILED


async def test_scan_during_a_datalab_outage_is_retryable_not_no_text(blank_pdf: Path) -> None:
    with pytest.raises(AppError) as caught:
        await _service(FakeConverter(fail=True)).parse(blank_pdf, ".pdf", external_parsing=True)
    assert caught.value.code is ErrorCode.PROVIDER_UNAVAILABLE and caught.value.retryable()


async def test_image_during_a_datalab_outage_is_retryable_not_no_text(tmp_path: Path) -> None:
    path = tmp_path / "a.png"
    path.write_bytes(b"\x89PNG\r\n\x1a\n")
    with pytest.raises(AppError) as caught:
        await _service(FakeConverter(fail=True)).parse(path, ".png", external_parsing=True)
    assert caught.value.code is ErrorCode.PROVIDER_UNAVAILABLE


async def test_empty_docx_during_a_datalab_outage_is_retryable(empty_docx: Path) -> None:
    with pytest.raises(AppError) as caught:
        await _service(FakeConverter(fail=True)).parse(empty_docx, ".docx", external_parsing=True)
    assert caught.value.code is ErrorCode.PROVIDER_UNAVAILABLE


async def test_no_text_stays_no_text_when_datalab_is_off_or_succeeded_empty(blank_pdf: Path) -> None:
    cases = [(None, True), (FakeConverter(), False), (FakeConverter(tree={"children": []}), True)]
    for converter, external in cases:
        with pytest.raises(AppError) as caught:
            await _service(converter).parse(blank_pdf, ".pdf", external_parsing=external)
        assert caught.value.code is ErrorCode.INGEST_NO_TEXT


async def test_span_attributes_land_on_every_parse_span(pdf: Path, docx: Path) -> None:
    exporter = InMemorySpanExporter()
    provider = TracerProvider()
    provider.add_span_processor(SimpleSpanProcessor(exporter))
    ids = {"request_id": "req-1", "org_id": "org-1", "kb_id": "kb-1", "file_id": "file-1"}
    service = _service(FakeConverter(qualities=[2.0, 3.0]), provider=provider)
    await service.parse(pdf, ".pdf", external_parsing=True, span_attributes=ids)
    await service.parse(docx, ".docx", external_parsing=True, span_attributes=ids)
    spans = exporter.get_finished_spans()
    assert {span.name for span in spans} == {SPAN_DATALAB, SPAN_LOCAL}
    assert len(spans) == 3  # two converts (quality retry) and one local parse
    for span in spans:
        assert all((span.attributes or {}).get(key) == value for key, value in ids.items())


async def test_datalab_span_records_the_cost_in_cents(pdf: Path) -> None:
    exporter = InMemorySpanExporter()
    provider = TracerProvider()
    provider.add_span_processor(SimpleSpanProcessor(exporter))
    converter = FakeConverter(cost={"final_cost_cents": 1.5})
    await _service(converter, provider=provider).parse(pdf, ".pdf", external_parsing=True)
    assert (_spans(exporter)[SPAN_DATALAB].attributes or {})["parse.cost_cents"] == 1.5


async def test_datalab_success_logs_its_duration(pdf: Path) -> None:
    with structlog.testing.capture_logs() as logs:
        await _service(FakeConverter()).parse(pdf, ".pdf", external_parsing=True)
    done = [entry for entry in logs if entry["event"] == "datalab convert completed"]
    assert len(done) == 1 and done[0]["mode"] == "balanced" and done[0]["duration_ms"] >= 0


async def test_datalab_error_logs_the_status_code_never_the_message(pdf: Path) -> None:
    with structlog.testing.capture_logs() as logs:
        await _service(FakeConverter(error=FakeApiError(503))).parse(pdf, ".pdf", external_parsing=True)
    failed = [entry for entry in logs if entry["event"] == "datalab failed, parsing locally"]
    assert failed[0]["status_code"] == 503 and failed[0]["error_type"] == "FakeApiError"
    assert "secret-looking text" not in str(logs)


async def test_a_failed_datalab_span_carries_the_error_type_never_the_message(pdf: Path) -> None:
    exporter = InMemorySpanExporter()
    provider = TracerProvider()
    provider.add_span_processor(SimpleSpanProcessor(exporter))
    await _service(FakeConverter(error=FakeApiError(503)), provider=provider).parse(
        pdf, ".pdf", external_parsing=True
    )
    span = _spans(exporter)[SPAN_DATALAB]
    assert span.status.status_code is StatusCode.ERROR
    assert span.status.description == "FakeApiError"
    assert "secret-looking text" not in str(span.events) + str(span.status.description)


class _LeakyParser:
    def parse(self, path: Path) -> ParsedDocument:
        raise ValueError("secret-looking text")


async def test_a_failed_local_span_carries_the_error_type_only(tmp_path: Path) -> None:
    path = tmp_path / "a.txt"
    path.write_text("x", encoding="utf-8")
    exporter = InMemorySpanExporter()
    provider = TracerProvider()
    provider.add_span_processor(SimpleSpanProcessor(exporter))
    service = ParseService(
        ParserRegistry({".txt": _LeakyParser()}),
        None,
        mode="balanced",
        retry_mode="accurate",
        min_quality=3.0,
        local_first_extensions=LOCAL_FIRST,
        tracer=provider.get_tracer(__name__),
    )
    with pytest.raises(AppError):
        await service.parse(path, ".txt", external_parsing=False)
    span = _spans(exporter)[SPAN_LOCAL]
    assert span.status.status_code is StatusCode.ERROR
    assert span.status.description == "ValueError"  # the parser's own type, never its message
    assert "secret-looking text" not in str(span.events) + str(span.attributes)
