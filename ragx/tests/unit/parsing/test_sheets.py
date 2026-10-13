from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from thomas_ragx.parsing.model import BlockType
from thomas_ragx.parsing.sheets import CSV_SHEET_NAME, CsvParser, XlsxParser


def test_csv_parser_yields_one_table_with_its_sheet_range(tmp_path: Path) -> None:
    path = tmp_path / "gia.csv"
    path.write_text("Gói,Giá\nA,10\nB,20\n", encoding="utf-8")
    (block,) = CsvParser().parse(path).blocks
    assert block.block_type is BlockType.TABLE
    # A constant, not the file stem: the pipeline parses a temp file named <file_id>-<version>.csv.
    assert block.locator.to_json() == {"kind": "sheet_range", "sheet": CSV_SHEET_NAME, "range": "A1:B3"}
    assert CSV_SHEET_NAME == "csv"
    assert block.table is not None and block.table.source_rows == (2, 3) and block.table.columns == ("A", "B")


def test_csv_sheet_name_does_not_depend_on_the_file_name(tmp_path: Path) -> None:
    path = tmp_path / "0b1c2d3e-4.csv"
    path.write_text("A,B\n1,2\n", encoding="utf-8")
    (block,) = CsvParser().parse(path).blocks
    assert block.locator.sheet == CSV_SHEET_NAME


def test_csv_table_width_stops_at_the_last_non_empty_column(tmp_path: Path) -> None:
    path = tmp_path / "wide.csv"
    path.write_text("a,b,,,\nc,d,,,\n", encoding="utf-8")
    (block,) = CsvParser().parse(path).blocks
    assert block.locator.cell_range == "A1:B2"
    assert block.table is not None and block.table.header == ("a", "b") and block.table.columns == ("A", "B")


def test_xlsx_styled_empty_cells_do_not_widen_the_table(tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.append(["Tên", "Số"])
    sheet.append(["An", 3])
    sheet["Z2"].font = Font(bold=True)  # a formatted but empty cell stretches the sheet dimension
    path = tmp_path / "styled.xlsx"
    book.save(path)
    table_block = XlsxParser().parse(path).blocks[1]
    assert table_block.locator.cell_range == "A1:B2"
    assert table_block.table is not None and table_block.table.rows == (("An", "3"),)


def test_xlsx_parser_gives_each_sheet_a_heading_and_a_table_and_skips_empty_rows(tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Quý 1"
    sheet.append(["Tên", "Số"])
    sheet.append([None, None])
    sheet.append(["An", 3])
    book.create_sheet("Trống")
    path = tmp_path / "b.xlsx"
    book.save(path)
    blocks = XlsxParser().parse(path).blocks
    assert [(b.block_type, b.text if b.block_type is BlockType.HEADING else None) for b in blocks] == [
        (BlockType.HEADING, "Quý 1"),
        (BlockType.TABLE, None),
    ]
    table = blocks[1].table
    assert table is not None and table.rows == (("An", "3"),) and table.source_rows == (3,)
    assert blocks[1].locator.cell_range == "A1:B3"
